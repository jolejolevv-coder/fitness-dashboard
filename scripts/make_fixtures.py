#!/usr/bin/env python3
"""Erzeugt Testfixtures aus einem lokalen Sheet Export.

Die Fixtures bilden nach, was die Google Sheets API v4 mit der Voreinstellung
valueRenderOption=FORMATTED_VALUE liefert: ein Array von Zeilen, jede Zeile ein
Array von Zeichenketten in der Darstellung des Sheets. Nachlaufende leere Zellen
und nachlaufende leere Zeilen laesst die API weg, das wird hier nachgebildet.

Der Rohexport liegt unter testdata/ und wird nie versioniert. Die erzeugten
Fixtures unter test/fixtures/ werden versioniert.

Aufruf:
    python3 scripts/make_fixtures.py <pfad-zum-xlsx>

Abhaengigkeit: openpyxl. Bewusst Python statt TypeScript, damit die
Laufzeitabhaengigkeiten des Projekts nicht um eine reine Entwicklungsbibliothek
zum Lesen von xlsx wachsen.
"""

import sys
import json
import datetime
from pathlib import Path

import openpyxl

# Bereiche exakt so, wie das Dashboard sie zur Laufzeit anfordert.
RANGES = {
    "daily-tracking": ("Daily Tracking", 300, 30),
    "mesocycle-1": ("Mesocycle 1", 200, 182),
    "mesocycle-2": ("Mesocycle 2", 200, 182),
}

# Spalten mit Freitext, 1 basiert. Inhalt wird durch Platzhalter ersetzt,
# damit keine persoenlichen Notizen in versionierte Testdaten geraten.
COMMENT_COLUMNS = {"Daily Tracking": (12, 29)}

PLACEHOLDER = "Notiz {n}"


def german_number(value):
    """Bildet die Automatik Darstellung von Google Sheets in deutscher Lokalisierung nach."""
    if isinstance(value, bool):
        return "WAHR" if value else "FALSCH"
    if float(value).is_integer():
        return str(int(value))
    # Sheets zeigt in der Automatik bis zu zehn signifikante Stellen.
    text = f"{value:.10g}"
    return text.replace(".", ",")


def german_percent(value, decimals=2):
    text = f"{value * 100:.{decimals}f}"
    return text.replace(".", ",") + "%"


def format_cell(cell):
    """Wandelt eine Zelle in die Zeichenkette, die die Sheets API liefern wuerde."""
    value = cell.value
    if value is None:
        return ""

    number_format = cell.number_format or "General"

    if isinstance(value, datetime.datetime):
        return f"{value.day}.{value.month}.{value.year}"
    if isinstance(value, datetime.date):
        return f"{value.day}.{value.month}.{value.year}"
    if isinstance(value, datetime.time):
        return f"{value.hour:02d}:{value.minute:02d}"

    if isinstance(value, (int, float)):
        if "%" in number_format:
            decimals = number_format.count("0", number_format.find(".") + 1) if "." in number_format else 0
            return german_percent(value, decimals)
        return german_number(value)

    return str(value).strip()


def trim_trailing(items):
    """Entfernt nachlaufende Leerwerte, wie es die API tut."""
    end = len(items)
    while end > 0 and (items[end - 1] == "" or items[end - 1] == []):
        end -= 1
    return items[:end]


def extract(worksheet, max_row, max_col, comment_columns):
    counter = 0
    rows = []
    for r in range(1, max_row + 1):
        row = []
        for c in range(1, max_col + 1):
            cell = worksheet.cell(r, c)
            text = format_cell(cell)
            if text and c in comment_columns:
                counter += 1
                text = PLACEHOLDER.format(n=counter)
            row.append(text)
        rows.append(trim_trailing(row))
    return trim_trailing(rows)


def main():
    if len(sys.argv) != 2:
        print(__doc__)
        return 1

    source = Path(sys.argv[1])
    if not source.exists():
        print(f"Export nicht gefunden: {source}")
        return 1

    target = Path(__file__).resolve().parent.parent / "test" / "fixtures"
    target.mkdir(parents=True, exist_ok=True)

    workbook = openpyxl.load_workbook(source, data_only=True)

    for name, (tab, max_row, max_col) in RANGES.items():
        if tab not in workbook.sheetnames:
            print(f"Tab fehlt, uebersprungen: {tab}")
            continue
        worksheet = workbook[tab]
        rows = extract(
            worksheet,
            min(max_row, worksheet.max_row),
            min(max_col, worksheet.max_column),
            set(COMMENT_COLUMNS.get(tab, ())),
        )
        out = target / f"{name}.json"
        out.write_text(json.dumps(rows, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"{out.name}: {len(rows)} Zeilen")

    return 0


if __name__ == "__main__":
    sys.exit(main())
